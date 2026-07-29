"""Generates slides/camunda8-order-fulfillment.pptx from the same content as order-fulfillment-slides.html.

Run: .venv/Scripts/python slides/build_pptx.py
Not part of the runtime lab -- a one-off deck builder, kept for easy re-runs
if the process diagram or slide content changes again.
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

INK = RGBColor(0x12, 0x15, 0x1B)
MUTED = RGBColor(0x56, 0x60, 0x73)
PAPER = RGBColor(0xEE, 0xF1, 0xF5)
PANEL = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xCF, 0xD6, 0xDF)
ACCENT = RGBColor(0xE8, 0x87, 0x1F)
ACCENT_INK = RGBColor(0x7A, 0x4A, 0x0C)
OK = RGBColor(0x2F, 0x8F, 0x5B)

FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
PROCESS_DIAGRAM = ASSETS_DIR / "process-diagram.png"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send background behind everything
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_text(slide, left, top, width, height, text, size, color=INK, bold=False,
             font=FONT_BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return box


def add_eyebrow(slide, text, top=Inches(0.55)):
    add_text(slide, Inches(0.7), top, Inches(8), Inches(0.35), text.upper(), 13,
              color=ACCENT, bold=True, font=FONT_MONO)


def add_title(slide, text, top=Inches(0.92), size=40, width=Inches(11.5)):
    add_text(slide, Inches(0.7), top, width, Inches(1.1), text, size, color=INK,
              bold=True, font=FONT_MONO)


def add_lede(slide, text, top, width=Inches(9)):
    add_text(slide, Inches(0.7), top, width, Inches(0.8), text, 16, color=MUTED, line_spacing=1.2)


def footer(slide, n, total=10):
    add_text(slide, Inches(0.7), Inches(7.1), Inches(4), Inches(0.3),
              "Camunda 8 · order-fulfillment demo", 10, color=MUTED, font=FONT_MONO)
    add_text(slide, Inches(12.2), Inches(7.1), Inches(0.7), Inches(0.3),
              f"{n} / {total}", 10, color=MUTED, font=FONT_MONO, align=PP_ALIGN.RIGHT)


def rect(slide, left, top, width, height, text, sub=None, accent=False, dashed=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.12
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = ACCENT if accent else LINE
    shp.line.width = Pt(1.75 if accent else 1)
    if dashed:
        ln = shp.line._get_or_add_ln()
        d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
        ln.append(d)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(13); r.font.color.rgb = INK; r.font.name = FONT_BODY; r.font.bold = True
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(9.5); r2.font.color.rgb = MUTED; r2.font.name = FONT_MONO
    return shp


def arrow(slide, x1, y1, x2, y2, label=None, label_top=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = MUTED
    conn.line.width = Pt(1.25)
    conn.line.end_arrowhead = True
    if label:
        add_text(slide, min(x1, x2), label_top or (min(y1, y2) - Inches(0.28)),
                  abs(x2 - x1) + Inches(1), Inches(0.25), label, 9.5, color=MUTED,
                  font=FONT_MONO, align=PP_ALIGN.CENTER)
    return conn


def circle(slide, cx, cy, r, filled=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    shp.fill.solid()
    shp.fill.fore_color.rgb = INK if filled else PANEL
    shp.line.color.rgb = INK
    shp.line.width = Pt(1.75)
    shp.shadow.inherit = False
    return shp


def gap_column(slide, left, header, header_color, items):
    add_text(slide, left, Inches(2.75), Inches(5.7), Inches(0.35), header, 13, color=header_color,
              bold=True, font=FONT_MONO)
    y = Inches(3.2)
    for head, body in items:
        box = slide.shapes.add_textbox(left, y, Inches(5.7), Inches(0.9))
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = head + " -- "; r1.font.bold = True; r1.font.size = Pt(13); r1.font.color.rgb = INK; r1.font.name = FONT_BODY
        r2 = p.add_run(); r2.text = body; r2.font.size = Pt(13); r2.font.color.rgb = MUTED; r2.font.name = FONT_BODY
        y += Inches(0.95)


# ---------------------------------------------------------------- Slide 1: Title
s = add_slide()
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit = False
add_eyebrow(s, "Live demo", top=Inches(2.3))
add_title(s, "Camunda 8,\nrunning for real.", top=Inches(2.7), size=54, width=Inches(11))
add_lede(s, "A hands-on exploration: one BPMN process, a Python worker, in/outbound connectors, "
             "a DMN decision, a form, and every boundary event type -- deployed, run, and broken live.",
         top=Inches(4.4), width=Inches(9.5))
add_text(s, Inches(0.7), Inches(6.3), Inches(9), Inches(0.4),
          "Zeebe · Operate · Tasklist  —  localhost:8080", 13, color=MUTED, font=FONT_MONO)

# ---------------------------------------------------------------- Slide 2: The problem
s = add_slide()
add_eyebrow(s, "Why orchestration")
add_title(s, "Business logic, scattered.")
bullets = [
    ("today", "Process steps live inside cron jobs, message handlers, and tribal knowledge -- no single place shows what a process actually does."),
    ("today", "When something is stuck, you grep logs across five services to find out why."),
    ("today", "Changing the order of steps means a code change, a review, and a deploy."),
]
y = Inches(2.2)
for tag, body in bullets:
    tagbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(1.0), Inches(0.4))
    tagbox.fill.background(); tagbox.line.color.rgb = LINE; tagbox.line.width = Pt(1); tagbox.shadow.inherit = False
    tf = tagbox.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = tag; r.font.size = Pt(11); r.font.color.rgb = ACCENT; r.font.name = FONT_MONO
    add_text(s, Inches(2.0), y, Inches(10.2), Inches(0.9), body, 15, color=INK, line_spacing=1.15,
              anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.1)
footer(s, 2)

# ---------------------------------------------------------------- Slide 3: Platform grid
s = add_slide()
add_eyebrow(s, "The platform")
add_title(s, "One diagram, executed.")
add_lede(s, "BPMN isn't documentation of the process -- it IS the process.", top=Inches(1.85))
cards = [
    ("BPMN", "The diagram. Defines steps, order, and decisions.", True),
    ("Zeebe", "The engine. Runs instances, scales horizontally.", True),
    ("DMN", "Decision tables, modeled separately. A business rule task calls one.", True),
    ("Connectors", "Prebuilt steps, configured not coded -- outbound (Charge Payment) and inbound (webhook).", False),
    ("Forms", "A JSON schema on a user task -- Tasklist renders it, no frontend code.", False),
    ("Boundary events", "Timer, error, and message -- three ways to interrupt a task.", False),
    ("Multi-instance", "Run one activity per element of a collection, not once per instance.", False),
    ("Operate", "Watch every instance live. See exactly where it's stuck.", False),
    ("Tasklist", "Where humans complete the manual steps.", False),
]
cols, rows = 3, 3
cw, ch = Inches(3.85), Inches(1.35)
gx, gy = Inches(0.25), Inches(0.25)
ox, oy = Inches(0.7), Inches(2.55)
for i, (title, body, core) in enumerate(cards):
    r, c = divmod(i, cols)
    left = ox + c * (cw + gx)
    top = oy + r * (ch + gy)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cw, ch)
    card.adjustments[0] = 0.06
    card.fill.solid(); card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = ACCENT if core else LINE
    card.line.width = Pt(1.75 if core else 1)
    card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(12); tf.margin_top = Pt(10); tf.margin_right = Pt(10)
    p = tf.paragraphs[0]
    r0 = p.add_run(); r0.text = title
    r0.font.size = Pt(15); r0.font.bold = True; r0.font.color.rgb = INK; r0.font.name = FONT_MONO
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = body
    r2.font.size = Pt(10.5); r2.font.color.rgb = MUTED; r2.font.name = FONT_BODY
footer(s, 3)

# ---------------------------------------------------------------- Slide 4: Architecture
s = add_slide()
add_eyebrow(s, "Under the hood")
add_title(s, "Five processes, three languages.")
add_lede(s, "None of them know the others' internals -- only the contracts.", top=Inches(1.85))

arch_y = Inches(2.6)
orchestration = rect(s, Inches(4.4), arch_y, Inches(4.5), Inches(0.85),
                      "orchestration", "Zeebe :26500 gRPC · Operate/Tasklist :8080", accent=True)

row2_y = arch_y + Inches(1.7)
order_service_box = rect(s, Inches(0.7), row2_y, Inches(3.2), Inches(0.8),
                          "order_service.py", "Python · FastAPI · :8000")
connectors_box = rect(s, Inches(9.5), row2_y, Inches(3.2), Inches(0.8),
                       "connectors", ":8086 · outbound + inbound webhook")

row3_y = row2_y + Inches(1.15)
order_workers_box = rect(s, Inches(0.7), row3_y, Inches(3.2), Inches(0.8),
                          "order_workers.py", "Python · pyzeebe · long-poll")
payment_box = rect(s, Inches(9.5), row3_y, Inches(3.2), Inches(0.8),
                    "payment_service.js", "Node · Express · :8001", accent=True)

arrow(s, Inches(3.9), row2_y + Inches(0.15), Inches(4.4), arch_y + Inches(0.55),
      "gRPC: start / cancel", row2_y - Inches(0.32))
arrow(s, Inches(3.9), row3_y + Inches(0.4), Inches(4.4), arch_y + Inches(0.75),
      "gRPC: poll / activate / complete", row3_y + Inches(0.85))
arrow(s, Inches(8.9), arch_y + Inches(0.55), Inches(9.5), row2_y + Inches(0.15),
      "gRPC job (built-in worker)", arch_y + Inches(0.9))
arrow(s, Inches(11.1), row2_y + Inches(0.8), Inches(11.1), row3_y,
      "HTTP POST", row2_y + Inches(0.82))
footer(s, 4)

# ---------------------------------------------------------------- Slide 5: Code vs configuration
s = add_slide()
add_eyebrow(s, "How a step runs")
add_title(s, "Code, or configuration.")
add_lede(s, "Three tasks in the same diagram, wired three different ways -- same engine, same contract.",
         top=Inches(1.85))

row_y = [Inches(2.55), Inches(4.05), Inches(5.55)]
row_label = ["JOB WORKER · Ship Order", "CONNECTOR · Charge Payment", "INBOUND CONNECTOR · Order Canceled"]
for label, y in zip(row_label, row_y):
    add_text(s, Inches(0.7), y, Inches(6), Inches(0.3), label, 11, color=MUTED, font=FONT_MONO)

y = row_y[0] + Inches(0.35)
rect(s, Inches(0.7), y, Inches(2.1), Inches(0.7), "Service Task", "type: ship-order", accent=True)
rect(s, Inches(4.3), y, Inches(2.5), Inches(0.7), "Python Worker", "order_workers.py")
arrow(s, Inches(2.8), y + Inches(0.22), Inches(4.3), y + Inches(0.22), "polls & activates", y - Inches(0.02))
arrow(s, Inches(4.3), y + Inches(0.48), Inches(2.8), y + Inches(0.48), "completes job", y + Inches(0.72))
add_text(s, Inches(7.1), y + Inches(0.22), Inches(2), Inches(0.3), "gRPC :26500", 10, color=MUTED, font=FONT_MONO)

y = row_y[1] + Inches(0.35)
rect(s, Inches(0.7), y, Inches(2.1), Inches(0.7), "Service Task", "type: http-json:1")
rect(s, Inches(3.4), y, Inches(2.3), Inches(0.7), "REST Connector", "built-in, zero code", accent=True)
rect(s, Inches(6.3), y, Inches(2.1), Inches(0.7), "payment_service", ":8001/charge")
arrow(s, Inches(2.8), y + Inches(0.35), Inches(3.4), y + Inches(0.35))
arrow(s, Inches(5.7), y + Inches(0.35), Inches(6.3), y + Inches(0.35), "HTTP POST", y + Inches(0.05))

y = row_y[2] + Inches(0.35)
rect(s, Inches(0.7), y, Inches(2.1), Inches(0.7), "order_service", "POST /cancel")
rect(s, Inches(3.4), y, Inches(2.4), Inches(0.7), "Webhook Connector", "/inbound/order-canceled", accent=True)
rect(s, Inches(6.5), y, Inches(2.3), Inches(0.7), "Boundary Event", "message, on Confirm Delivery")
arrow(s, Inches(2.8), y + Inches(0.35), Inches(3.4), y + Inches(0.35), "HTTP POST", y + Inches(0.05))
arrow(s, Inches(5.8), y + Inches(0.35), Inches(6.5), y + Inches(0.35), "correlates by orderId", y + Inches(0.05))
footer(s, 5)

# ---------------------------------------------------------------- Slide 6: The process (hero)
s = add_slide()
add_eyebrow(s, "Today's process")
add_title(s, "Order Fulfillment")
add_lede(s, "Same happy path as before -- now with six ways an instance can actually end.",
         top=Inches(1.85))
if PROCESS_DIAGRAM.exists():
    pic = s.shapes.add_picture(str(PROCESS_DIAGRAM), Inches(0.5), Inches(2.6), width=Inches(12.3))
footer(s, 6)

# ---------------------------------------------------------------- Slide 7: Failure demo
s = add_slide()
add_eyebrow(s, "When it breaks")
add_title(s, "No crash. No redeploy.")
add_lede(s, "Bad input reaches the worker -- instead of losing the instance, Zeebe holds it exactly where it failed.",
         top=Inches(1.85))
code = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(2.6), Inches(8.5), Inches(0.5))
code.fill.solid(); code.fill.fore_color.rgb = PANEL
code.line.color.rgb = LINE; code.line.width = Pt(1)
code.shadow.inherit = False
tf = code.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Pt(14)
p = tf.paragraphs[0]; r = p.add_run()
r.text = 'POST /orders {"order_id":"ORD-1003","quantity":0}'
r.font.size = Pt(13); r.font.name = FONT_MONO; r.font.color.rgb = INK

incident = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(3.35), Inches(8.5), Inches(1.0))
incident.adjustments[0] = 0.08
incident.fill.solid(); incident.fill.fore_color.rgb = PANEL
incident.line.color.rgb = RGBColor(0xC0, 0x43, 0x3A); incident.line.width = Pt(1.25)
incident.shadow.inherit = False
tf = incident.text_frame; tf.margin_left = Pt(16); tf.margin_top = Pt(10); tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
r = p.add_run(); r.text = "INCIDENT · Task_ValidateOrder"
r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x43, 0x3A); r.font.name = FONT_MONO
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.LEFT
r2 = p2.add_run()
r2.text = "Failed job. Error: Invalid quantity 0 for order ORD-1003"
r2.font.size = Pt(12); r2.font.name = FONT_MONO; r2.font.color.rgb = INK

add_text(s, Inches(0.7), Inches(4.7), Inches(5.6), Inches(0.35), "WITHOUT ORCHESTRATION", 11,
          color=RGBColor(0xC0, 0x43, 0x3A), bold=True, font=FONT_MONO)
add_text(s, Inches(0.7), Inches(5.1), Inches(5.6), Inches(1.2),
          "A failed script exits. Someone notices a missing order, hours later, from a support ticket.",
          14, color=MUTED, line_spacing=1.2)
add_text(s, Inches(6.8), Inches(4.7), Inches(5.6), Inches(0.35), "WITH ZEEBE", 11,
          color=OK, bold=True, font=FONT_MONO)
add_text(s, Inches(6.8), Inches(5.1), Inches(5.6), Inches(1.2),
          "Operate shows the exact instance and error. Fix the variable, click retry, the instance "
          "continues -- mid-flow.", 14, color=MUTED, line_spacing=1.2)
footer(s, 7)

# ---------------------------------------------------------------- Slide 8: Why it matters
s = add_slide()
add_eyebrow(s, "Why it matters")
add_title(s, "What you actually get.")
points = [
    ("see", "Every instance, live.", "We forced a real incident earlier -- Operate showed the exact failed variable and a one-click retry, not a log line to go reconstruct after the fact."),
    ("change", "Reroute without a deploy.", "The in-stock threshold is a row in process/dmn/stock-check.dmn, not a Python if -- change the number, redeploy the table, zero code review."),
    ("mix", "Any language, per step.", "The order-triggering service is Python; the payment service the connector calls is Node -- proven today, not hypothetical."),
    ("trust", "Nothing silently drops.", "Two different failure modes, two different outcomes: an unhandled bug becomes a visible incident; a declined payment becomes a modeled path, not a crash. We triggered both."),
]
y = Inches(2.3)
for tag, head, body in points:
    tagbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(1.1), Inches(0.4))
    tagbox.fill.background(); tagbox.line.color.rgb = LINE; tagbox.line.width = Pt(1); tagbox.shadow.inherit = False
    tf = tagbox.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = tag; r.font.size = Pt(11); r.font.color.rgb = ACCENT; r.font.name = FONT_MONO
    box = s.shapes.add_textbox(Inches(2.1), y - Inches(0.05), Inches(10.2), Inches(1.0))
    tf2 = box.text_frame; tf2.word_wrap = True
    p1 = tf2.paragraphs[0]
    r1 = p1.add_run(); r1.text = head + "  "; r1.font.bold = True; r1.font.size = Pt(15); r1.font.color.rgb = INK; r1.font.name = FONT_BODY
    r2 = p1.add_run(); r2.text = body; r2.font.size = Pt(15); r2.font.color.rgb = MUTED; r2.font.name = FONT_BODY
    y += Inches(1.15)
footer(s, 8)

# ---------------------------------------------------------------- Slide 9: What this doesn't prove yet
s = add_slide()
add_eyebrow(s, "Before anyone asks")
add_title(s, "What this doesn't prove yet.")
add_lede(s, "Breadth of BPMN concepts is covered. Two different kinds of gap remain, and they're not the same kind of problem.",
         top=Inches(1.85))
gap_column(s, Inches(0.7), "FEATURE GAPS", MUTED, [
    ("Parallel / inclusive gateways", "only the exclusive gateway exists today."),
    ("Compensation events", "Cancel Order doesn't roll back a reservation or a charge."),
    ("Process versioning & migration", "deploying v2 of a running process, live."),
])
gap_column(s, Inches(6.9), "PRODUCTION GAPS", ACCENT, [
    ("Auth", "still the default demo/demo login; Identity untouched."),
    ("Storage", "H2, single-node, file-based; production needs the Elasticsearch exporter or SaaS."),
    ("Observability & HA", "Operate's UI only, no alerting; one broker container, not a cluster."),
])
add_text(s, Inches(0.7), Inches(6.6), Inches(11.5), Inches(0.7),
          'The left column is "haven\'t gotten to it yet." The right column is "deliberately out of '
          'scope for a local demo" -- worth being explicit about which is which.',
          12, color=MUTED, line_spacing=1.2)
footer(s, 9)

# ---------------------------------------------------------------- Slide 10: Close
s = add_slide()
add_eyebrow(s, "Thanks", top=Inches(2.6))
add_title(s, "Questions?", top=Inches(3.0), size=54, width=Inches(9))
add_lede(s, "Repo, BPMN file, and worker code are all right here -- happy to walk through any piece live.",
         top=Inches(4.35), width=Inches(9))
add_text(s, Inches(0.7), Inches(5.3), Inches(9), Inches(0.4),
          "camunda8-lab / bpmn / dmn / forms / workers / services_python / services_node", 13, color=MUTED, font=FONT_MONO)

out = Path(__file__).resolve().parent / "camunda8-order-fulfillment.pptx"
prs.save(str(out))
print(f"Saved {out}")
